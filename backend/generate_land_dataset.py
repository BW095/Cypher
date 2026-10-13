import os
import shutil
from pathlib import Path
from docx import Document
from docx.shared import Inches
from pptx import Presentation
from pptx.util import Inches as PptxInches
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from PIL import Image, ImageDraw, ImageFont
import time

DATASET_ROOT = Path(__file__).parent.parent / "demo_dataset_land"

def main():
    if DATASET_ROOT.exists():
        shutil.rmtree(DATASET_ROOT)
    DATASET_ROOT.mkdir(parents=True)
    print(f"\n[BUILD] Building Land Records dataset at: {DATASET_ROOT}\n")

    # 1. Text Files
    txt_dir = DATASET_ROOT / "Text_Records"
    txt_dir.mkdir(parents=True)
    write_txt(txt_dir / "7_12_Extract_Survey_145_Pune.txt", "GOVERNMENT OF MAHARASHTRA\nREVENUE DEPARTMENT\n\n7/12 EXTRACT\nVillage: Kothrud, Taluka: Haveli, District: Pune\nSurvey No: 145/2/B\nArea: 2.5 Hectares\nOwner: Ramesh Kumar Deshmukh\nCrop: Sugarcane\nLoans: State Bank of India - Rs 5,00,000 against this plot.")
    write_txt(txt_dir / "Mutation_Entry_2023_045.txt", "MUTATION REGISTER ENTRY\nEntry No: ME-2023-045\nDate: 15-Sep-2023\nNature of Right: Inheritance\nFrom: Late Suresh Deshmukh\nTo: Ramesh Kumar Deshmukh (Son)\nApproved By: Circle Officer, Haveli.")
    print("  [OK] Text files created.")

    # 2. DOCX Files
    doc_dir = DATASET_ROOT / "Word_Documents"
    doc_dir.mkdir(parents=True)
    
    doc = Document()
    doc.add_heading('Land Revenue Policy Framework 2024', 0)
    doc.add_paragraph('This document outlines the revised guidelines for property tax assessment in urban areas.')
    doc.add_paragraph('Section 1: Commercial Properties will see a 5% increase in base tax rate.')
    doc.save(doc_dir / "Revenue_Policy_2024.docx")
    
    # Docx with image
    doc2 = Document()
    doc2.add_heading('Land Survey Report - Plot 89', 0)
    doc2.add_paragraph('Survey Conducted on: 12-Oct-2023\nSurveyor: Mahesh Patil\nFindings: Encroachment of 200 sq.ft on eastern boundary by adjacent shop.')
    
    # Create a dummy image to insert
    img_path = str(DATASET_ROOT / "temp_map.png")
    create_dummy_image(img_path, "Survey Map Area 89", text_color="red")
    doc2.add_picture(img_path, width=Inches(4))
    doc2.save(doc_dir / "Survey_Report_Plot89.docx")
    os.remove(img_path)
    print("  [OK] Word documents created.")

    # 3. PPTX Files
    ppt_dir = DATASET_ROOT / "Presentations"
    ppt_dir.mkdir(parents=True)
    prs = Presentation()
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Digitization of Land Records"
    subtitle.text = "Q3 Progress Report - Revenue Dept"
    
    slide_layout2 = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(slide_layout2)
    title2 = slide2.shapes.title
    title2.text = "Key Milestones"
    body = slide2.shapes.placeholders[1]
    tf = body.text_frame
    tf.text = "Phase 1: Scanning completed\nPhase 2: Data entry in progress\nPhase 3: Blockchain integration planned"
    prs.save(ppt_dir / "Digitization_Progress_Q3.pptx")
    print("  [OK] PowerPoint presentation created.")

    # 4. PDF Files (Text & Scanned Mock)
    pdf_dir = DATASET_ROOT / "PDF_Records"
    pdf_dir.mkdir(parents=True)
    
    c = canvas.Canvas(str(pdf_dir / "Court_Order_Land_Dispute_451.pdf"), pagesize=letter)
    c.drawString(100, 750, "CIVIL COURT, PUNE DISTRICT")
    c.drawString(100, 730, "Order No: CIVIL/2024/451")
    c.drawString(100, 700, "In the matter of Ramesh Kumar vs. State of Maharashtra.")
    c.drawString(100, 680, "Order: The stay on construction at Survey No 145/2/B is hereby lifted.")
    c.save()

    # Create a 'scanned' looking PDF by generating an image and saving as PDF
    scan_img_path = str(DATASET_ROOT / "temp_scan.jpg")
    create_dummy_image(scan_img_path, "SALE DEED - SCANNED COPY\nSeller: Anil Sharma\nBuyer: Ramesh Kumar\nConsideration: Rs 45,00,000", bg_color="#f0f0f0", noise=True)
    scan_img = Image.open(scan_img_path)
    scan_img.save(str(pdf_dir / "Scanned_Sale_Deed_1998.pdf"), "PDF", resolution=100.0)
    os.remove(scan_img_path)
    print("  [OK] PDF records created.")

    # 5. Images
    img_dir = DATASET_ROOT / "Images"
    img_dir.mkdir(parents=True)
    
    # Image with text (like a photo of a document)
    create_dummy_image(str(img_dir / "Photo_Property_Tax_Receipt.jpg"), "PROPERTY TAX RECEIPT\nAmount Paid: Rs 12,500\nDate: 05-Mar-2024\nTransaction ID: TXN998822")
    
    # Image without text (just a colored block to simulate a drone photo or map)
    create_dummy_image(str(img_dir / "Drone_Survey_Plot145.jpg"), "", bg_color="green", noise=False)
    print("  [OK] Image files created.")

    print("\n" + "="*60)
    print("  LAND RECORDS DATASET CREATED SUCCESSFULLY!")
    print("="*60)
    print(f"\n  Location: {DATASET_ROOT}")

def write_txt(path: Path, content: str):
    path.write_text(content.strip(), encoding="utf-8")

def create_dummy_image(path: str, text: str, text_color="black", bg_color="white", noise=False):
    img = Image.new('RGB', (800, 600), color = bg_color)
    d = ImageDraw.Draw(img)
    if noise:
        import random
        for i in range(1000):
            d.point((random.randint(0, 800), random.randint(0, 600)), fill="gray")
    if text:
        d.text((50, 50), text, fill=text_color)
    img.save(path)

if __name__ == "__main__":
    main()
